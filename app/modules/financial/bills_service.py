import io
import calendar
from datetime import date
from typing import Optional, List, Tuple
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.modules.financial.models import FinancialBill
from app.modules.financial.bills_schemas import (
    FinancialBillCreate,
    FinancialBillUpdate,
    FinancialBillSettle,
    FinancialBillResponse,
    FinancialBillsSummaryResponse,
    FinancialBillListResponse,
)
from app.modules.financial.bills_repository import FinancialBillsRepository


def add_months(sourcedate: date, months: int) -> date:
    """Incrementa N meses mantendo o dia ou ajustando ao último dia do mês."""
    month = sourcedate.month - 1 + months
    year = sourcedate.year + month // 12
    month = month % 12 + 1
    day = min(sourcedate.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


class FinancialBillsService:
    def __init__(self):
        self.repo = FinancialBillsRepository()

    def _to_response(self, bill: FinancialBill) -> FinancialBillResponse:
        return FinancialBillResponse(
            id=bill.id,
            tenant_id=bill.tenant_id,
            bill_type=bill.bill_type,
            description=bill.description,
            category_id=bill.category_id,
            category_name=bill.category.name if bill.category else None,
            supplier_id=bill.supplier_id,
            supplier_name=bill.supplier.name if bill.supplier else None,
            client_id=bill.client_id,
            client_name=bill.client.name if bill.client else None,
            sale_id=bill.sale_id,
            amount=float(bill.amount or 0.0),
            paid_amount=float(bill.paid_amount or 0.0),
            discount_amount=float(bill.discount_amount or 0.0),
            fine_or_interest_amount=float(bill.fine_or_interest_amount or 0.0),
            issue_date=bill.issue_date,
            due_date=bill.due_date,
            payment_date=bill.payment_date,
            status=bill.status,
            payment_method=bill.payment_method,
            document_number=bill.document_number,
            barcode=bill.barcode,
            installment_number=bill.installment_number or 1,
            total_installments=bill.total_installments or 1,
            parent_bill_id=bill.parent_bill_id,
            destination_account_id=bill.destination_account_id,
            notes=bill.notes,
            created_at=bill.created_at,
            updated_at=bill.updated_at,
        )

    def create_bill(
        self, db: Session, tenant_id: int, data: FinancialBillCreate, user_id: Optional[int] = None
    ) -> List[FinancialBillResponse]:
        total_installments = data.total_installments or 1
        issue_date = data.issue_date or date.today()

        if total_installments == 1:
            bill = FinancialBill(
                tenant_id=tenant_id,
                bill_type=data.bill_type,
                description=data.description,
                category_id=data.category_id,
                supplier_id=data.supplier_id if data.bill_type == "payable" else None,
                client_id=data.client_id if data.bill_type == "receivable" else None,
                amount=data.amount,
                issue_date=issue_date,
                due_date=data.due_date,
                payment_method=data.payment_method,
                document_number=data.document_number,
                barcode=data.barcode,
                notes=data.notes,
                installment_number=1,
                total_installments=1,
                created_by_user_id=user_id,
            )
            created = self.repo.create_bill(db, bill)
            return [self._to_response(created)]

        # Criação parcelada
        total_amount = round(data.amount, 2)
        base_amount = round(total_amount / total_installments, 2)
        remainder = round(total_amount - (base_amount * total_installments), 2)

        # Cria a primeira parcela como mãe
        first_bill = FinancialBill(
            tenant_id=tenant_id,
            bill_type=data.bill_type,
            description=f"{data.description} (1/{total_installments})",
            category_id=data.category_id,
            supplier_id=data.supplier_id if data.bill_type == "payable" else None,
            client_id=data.client_id if data.bill_type == "receivable" else None,
            amount=round(base_amount + remainder, 2),  # ajusta centavos na primeira
            issue_date=issue_date,
            due_date=data.due_date,
            payment_method=data.payment_method,
            document_number=f"{data.document_number}-1" if data.document_number else None,
            barcode=data.barcode,
            notes=data.notes,
            installment_number=1,
            total_installments=total_installments,
            created_by_user_id=user_id,
        )
        first_bill = self.repo.create_bill(db, first_bill)
        bills_list = [first_bill]

        # Cria as parcelas subsequentes
        for i in range(2, total_installments + 1):
            next_due_date = add_months(data.due_date, i - 1)
            next_bill = FinancialBill(
                tenant_id=tenant_id,
                bill_type=data.bill_type,
                description=f"{data.description} ({i}/{total_installments})",
                category_id=data.category_id,
                supplier_id=data.supplier_id if data.bill_type == "payable" else None,
                client_id=data.client_id if data.bill_type == "receivable" else None,
                amount=base_amount,
                issue_date=issue_date,
                due_date=next_due_date,
                payment_method=data.payment_method,
                document_number=f"{data.document_number}-{i}" if data.document_number else None,
                barcode=data.barcode,
                notes=data.notes,
                installment_number=i,
                total_installments=total_installments,
                parent_bill_id=first_bill.id,
                created_by_user_id=user_id,
            )
            bills_list.append(next_bill)

        created_batch = self.repo.create_bills_batch(db, bills_list[1:])
        all_created = [first_bill] + created_batch
        return [self._to_response(b) for b in all_created]

    def get_bills(
        self,
        db: Session,
        tenant_id: int,
        bill_type: Optional[str] = None,
        status: Optional[str] = None,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        category_id: Optional[int] = None,
        supplier_id: Optional[int] = None,
        client_id: Optional[int] = None,
        search: Optional[str] = None,
        page: int = 1,
        page_size: int = 50,
        sort_by: str = "due_date",
        sort_dir: str = "asc",
    ) -> FinancialBillListResponse:
        items, total = self.repo.get_bills(
            db=db,
            tenant_id=tenant_id,
            bill_type=bill_type,
            status=status,
            start_date=start_date,
            end_date=end_date,
            category_id=category_id,
            supplier_id=supplier_id,
            client_id=client_id,
            search=search,
            page=page,
            page_size=page_size,
            sort_by=sort_by,
            sort_dir=sort_dir,
        )
        total_pages = (total + page_size - 1) // page_size if page_size > 0 else 1
        return FinancialBillListResponse(
            items=[self._to_response(item) for item in items],
            total=total,
            page=page,
            page_size=page_size,
            total_pages=total_pages,
        )

    def get_summary(
        self,
        db: Session,
        tenant_id: int,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        category_id: Optional[int] = None,
    ) -> FinancialBillsSummaryResponse:
        return self.repo.get_summary(
            db=db,
            tenant_id=tenant_id,
            start_date=start_date,
            end_date=end_date,
            category_id=category_id,
        )

    def get_bill(self, db: Session, tenant_id: int, bill_id: int) -> Optional[FinancialBillResponse]:
        bill = self.repo.get_bill(db, tenant_id=tenant_id, bill_id=bill_id)
        if not bill:
            return None
        return self._to_response(bill)

    def update_bill(
        self, db: Session, tenant_id: int, bill_id: int, data: FinancialBillUpdate
    ) -> Optional[FinancialBillResponse]:
        bill = self.repo.update_bill(db, tenant_id=tenant_id, bill_id=bill_id, data=data)
        if not bill:
            return None
        return self._to_response(bill)

    def settle_bill(
        self, db: Session, tenant_id: int, bill_id: int, data: FinancialBillSettle
    ) -> Optional[FinancialBillResponse]:
        bill = self.repo.settle_bill(db, tenant_id=tenant_id, bill_id=bill_id, data=data)
        if not bill:
            return None
        return self._to_response(bill)

    def cancel_bill(self, db: Session, tenant_id: int, bill_id: int) -> Optional[FinancialBillResponse]:
        bill = self.repo.cancel_bill(db, tenant_id=tenant_id, bill_id=bill_id)
        if not bill:
            return None
        return self._to_response(bill)

    def delete_bill(self, db: Session, tenant_id: int, bill_id: int) -> bool:
        return self.repo.delete_bill(db, tenant_id=tenant_id, bill_id=bill_id)

    def export_bills_excel(
        self,
        db: Session,
        tenant_id: int,
        bill_type: Optional[str] = None,
        status: Optional[str] = None,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        category_id: Optional[int] = None,
    ) -> io.BytesIO:
        bills = self.repo.get_all_for_export(
            db=db,
            tenant_id=tenant_id,
            bill_type=bill_type,
            status=status,
            start_date=start_date,
            end_date=end_date,
            category_id=category_id,
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Contas a Pagar e Receber"
        ws.views.sheetView[0].showGridLines = True

        # Paleta de estilos
        primary_color = "313485"
        header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        # Cabeçalho Principal
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "RELATÓRIO DE CONTAS A PAGAR E RECEBER"
        title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        title_cell.fill = header_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        headers = [
            "Tipo",
            "Descrição",
            "Categoria (DRE)",
            "Fornecedor / Cliente",
            "Vencimento",
            "Emissão",
            "Valor Nominal",
            "Valor Pago",
            "Pagamento",
            "Status",
            "Nº Documento",
        ]

        ws.append([])  # Linha 2 vazia
        ws.append(headers)  # Linha 3 cabeçalhos
        ws.row_dimensions[3].height = 24

        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Preenchimento de linhas
        status_labels = {
            "pending": "A Vencer",
            "paid": "Liquidado",
            "overdue": "Vencido",
            "canceled": "Cancelado",
        }

        row_num = 4
        for b in bills:
            contact = b.supplier.name if b.supplier else (b.client.name if b.client else "-")
            category = b.category.name if b.category else "-"
            type_label = "A Pagar" if b.bill_type == "payable" else "A Receber"
            status_text = status_labels.get(b.status, b.status)

            due_str = b.due_date.strftime("%d/%m/%Y") if b.due_date else "-"
            issue_str = b.issue_date.strftime("%d/%m/%Y") if b.issue_date else "-"
            pay_str = b.payment_date.strftime("%d/%m/%Y") if b.payment_date else "-"

            row_data = [
                type_label,
                b.description,
                category,
                contact,
                due_str,
                issue_str,
                float(b.amount or 0.0),
                float(b.paid_amount or 0.0),
                pay_str,
                status_text,
                b.document_number or "-",
            ]
            ws.append(row_data)

            # Estilização da linha
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = thin_border
                cell.font = Font(name="Arial", size=9)

                # Alinhamentos e formatações
                if col_idx in (5, 6, 9, 10):  # Datas e Status
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in (7, 8):  # Moeda
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '"R$" #,##0.00'
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            row_num += 1

        # Ajuste de largura das colunas
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        ws.column_dimensions["B"].width = 30  # Descrição
        ws.column_dimensions["C"].width = 22  # Categoria
        ws.column_dimensions["D"].width = 25  # Contato

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
